from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "KM3_Raport_Finalny.docx"
OUT_PPTX = ROOT / "KM3_Prezentacja.pptx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def build_report() -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Raport finalny KM3\n")
    run.bold = True
    run.font.size = Pt(20)
    subtitle = title.add_run("Climate, Economy and Mental Health in Europe (2000-2010)")
    subtitle.font.size = Pt(14)
    doc.add_paragraph("Zespół: Liwia Jankowska, Wiktoria Kawa, Wojciech Łukaszyk").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Hurtownie danych i systemy Business Intelligence").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_heading(doc, "1. Cel systemu i korzyści biznesowe")
    doc.add_paragraph(
        "Celem projektu jest budowa systemu klasy DW/BI integrującego dane klimatyczne, "
        "ekonomiczne oraz dane o rozpowszechnieniu wybranych zaburzeń psychicznych dla krajów "
        "europejskich w latach 2000-2010. System ma wspierać analizę zależności pomiędzy "
        "warunkami klimatycznymi, czynnikami ekonomicznymi i wskaźnikami zdrowia psychicznego."
    )
    add_bullets(
        doc,
        [
            "identyfikacja krajów i regionów o podwyższonej prevalencji wybranych zaburzeń psychicznych;",
            "analiza zależności pomiędzy zdrowiem psychicznym, bezrobociem, PKB i spożyciem alkoholu;",
            "analiza roli czynników klimatycznych, w szczególności nasłonecznienia i temperatury zimowej;",
            "udostępnienie interaktywnych raportów Power BI dla analityka zdrowia publicznego lub decydenta.",
        ],
    )

    add_heading(doc, "2. Finalna architektura rozwiązania")
    doc.add_paragraph(
        "Finalna architektura została zrealizowana jako warstwowy proces ETL/ELT z oddzieleniem "
        "źródeł danych, warstwy staging, hurtowni danych oraz raportowania BI."
    )
    add_table(
        doc,
        ["Warstwa", "Technologia / komponent", "Rola"],
        [
            ["Źródła", "Open-Meteo API, lokalne CSV", "Dane pogodowe, ekonomiczne i zdrowia psychicznego."],
            ["ETL", "Python", "Pobieranie danych, mapowanie kodów, obsługa CSV, cache API, zapis do staging."],
            ["Staging", "SQL Server, schema staging", "Przechowanie danych po ekstrakcji przed integracją."],
            ["ELT / DW", "SQL Server, schema dw", "Transformacje, wymiary, fakty, agregaty i widoki analityczne."],
            ["BI", "Power BI Desktop", "Model semantyczny, miary, interaktywne raporty i eksport PDF."],
        ],
    )
    doc.add_paragraph(
        "W raporcie KM2 rozważano użycie narzędzi SSIS, ADF lub dbt. W finalnej implementacji "
        "zastosowano moduł Python jako warstwę ETL, natomiast transformacje hurtowniane wykonano "
        "w SQL Serverze. Decyzja ta ułatwiła integrację z API Open-Meteo, obsługę lokalnych plików CSV "
        "oraz implementację cache i ponawiania zapytań w przypadku limitów API."
    )

    add_heading(doc, "3. Źródła danych")
    add_table(
        doc,
        ["Źródło", "Zakres", "Format", "Odświeżanie / dostęp", "Wykorzystanie"],
        [
            ["Open-Meteo Historical API", "Dane dzienne 2000-2010", "JSON API", "Dostęp przez HTTP; zastosowano lokalny cache", "Temperatura, opady, nasłonecznienie, zachmurzenie, wiatr."],
            ["FactEconomyYearly.csv", "Dane roczne 2000-2010", "CSV", "Lokalny plik źródłowy", "PKB per capita, wydatki zdrowotne, bezrobocie, alkohol."],
            ["FactMentalHealthYearly.csv", "Dane roczne 2000-2010", "CSV", "Lokalny plik źródłowy", "Prevalencja zaburzeń psychicznych."],
            ["countries_europe.json", "Metadane krajów", "JSON", "Konfiguracja ETL", "Mapowanie ISO2/ISO3, region, income group, współrzędne stolic."],
        ],
    )

    add_heading(doc, "4. Proces ETL/ELT")
    doc.add_paragraph(
        "Proces ETL uruchamiany jest poleceniem python run_etl.py. W pierwszym kroku Python czyści "
        "warstwę staging, pobiera lub odczytuje dane źródłowe, wykonuje podstawowe mapowania i zapisuje "
        "dane do tabel stagingowych. Następnie uruchamiany jest skrypt SQL load_dw_from_staging.sql, "
        "który ładuje dane do wymiarów i faktów hurtowni."
    )
    add_table(
        doc,
        ["Transformacja", "Opis"],
        [
            ["ISO3 -> ISO2", "Pliki ekonomiczne używają kodów ISO3, natomiast hurtownia używa kodów ISO2."],
            ["Mapowanie zaburzeń", "Szczegółowe jednostki, np. Major depressive disorder i Persistent depressive disorder, są agregowane do Depressive disorders."],
            ["Nasłonecznienie", "Open-Meteo zwraca czas nasłonecznienia w sekundach; w DW jest przeliczany na godziny."],
            ["Grain alignment", "Pogoda jest dzienna, a ekonomia i zdrowie psychiczne roczne; dodano widoki agregujące pogodę miesięcznie, rocznie i sezonowo."],
            ["Cache API", "Odpowiedzi Open-Meteo są zapisywane lokalnie, aby uniknąć ponownego pobierania tych samych danych i ograniczyć skutki limitów API."],
            ["Ponowne ładowanie", "Przed załadowaniem faktów usuwane są rekordy dla przetwarzanego ziarna, co zapobiega duplikatom."],
        ],
    )

    add_heading(doc, "5. Model fizyczny hurtowni danych")
    doc.add_paragraph(
        "Model hurtowni zrealizowano jako schemat konstelacji z trzema tabelami faktów współdzielącymi "
        "wymiary kraju i czasu. Dane są przechowywane w dwóch schematach: staging oraz dw."
    )
    add_table(
        doc,
        ["Typ", "Obiekty"],
        [
            ["Staging", "STG_COUNTRY, STG_WEATHER, STG_ECONOMY, STG_MENTAL_HEALTH"],
            ["Wymiary", "DIM_COUNTRY, DIM_DATE, DIM_YEAR, DIM_MONTH, DIM_DISORDER"],
            ["Fakty", "FACT_WEATHER, FACT_ECONOMY, FACT_MENTAL_HEALTH"],
            ["Widoki agregujące", "AGG_WEATHER_MONTHLY, AGG_WEATHER_ANNUAL, AGG_WEATHER_SEASONAL"],
        ],
    )
    add_table(
        doc,
        ["Tabela faktów", "Ziarno", "Najważniejsze miary"],
        [
            ["FACT_WEATHER", "kraj x dzień", "TEMP_AVG, TEMP_MIN, TEMP_MAX, PRECIP_MM, SUNSHINE_H, CLOUD_COVER, WIND_SPEED"],
            ["FACT_ECONOMY", "kraj x rok", "GDP_PER_CAPITA_USD, HEALTH_EXP_PER_CAPITA_USD, ALCOHOL_LITERS_PER_CAPITA, UNEMPLOYMENT_PCT"],
            ["FACT_MENTAL_HEALTH", "kraj x rok x zaburzenie", "PREVALENCE_PCT, DALYS_PER_100K"],
        ],
    )

    add_heading(doc, "6. Miary, atrybuty i hierarchie")
    add_bullets(
        doc,
        [
            "Hierarchia czasu: rok, kwartał, miesiąc, dzień; w Power BI użyto przede wszystkim DIM_YEAR i DIM_MONTH.",
            "Hierarchia geograficzna: region -> kraj.",
            "Hierarchia zaburzeń: kategoria -> nazwa zaburzenia.",
            "Miary BI: średnia prevalencja, średni PKB per capita, średnie bezrobocie, średnie spożycie alkoholu, średnia temperatura, średnie nasłonecznienie, zimowe nasłonecznienie.",
        ],
    )

    add_heading(doc, "7. Warstwa raportowa Power BI")
    doc.add_paragraph(
        "Warstwa raportowa została wykonana w Power BI Desktop w trybie Import. Do modelu załadowano "
        "wymiary, tabele faktów rocznych oraz widoki agregujące pogodę. Dane dzienne pogodowe nie muszą "
        "być bezpośrednio analizowane w raportach, ponieważ ich odpowiedniki miesięczne, roczne i sezonowe "
        "są dostępne w widokach agregujących."
    )
    add_table(
        doc,
        ["Strona raportu", "Cel", "Najważniejsze wizualizacje"],
        [
            ["Mental Health Map", "Analiza przestrzenna prevalencji zaburzeń.", "Kartogram/mapa, filtry regionu, income group, roku i zaburzenia."],
            ["Mental Health & Economy", "Porównanie zdrowia psychicznego z czynnikami ekonomicznymi.", "Scatter plot GDP vs prevalence, trendy roczne, tabela wskaźników."],
            ["Climate & Mental Health", "Analiza związku klimatu i zdrowia psychicznego.", "Scatter winter sunshine vs prevalence, ranking zimowego nasłonecznienia, temperatura miesięczna."],
            ["Alcohol & Unemployment", "Analiza alkoholu, bezrobocia i zaburzeń związanych z alkoholem.", "Scatter unemployment vs alcohol, ranking krajów, filtry regionu i income group."],
            ["What-if Simulation", "Symulacja scenariusza zmiany temperatury zimowej.", "Parametr what-if, baseline vs simulated depression prevalence."],
        ],
    )
    doc.add_paragraph("Załącznik: Power_BI_Raport.pdf zawiera wyeksportowane strony raportu Power BI.")

    add_heading(doc, "8. Analiza jakości danych")
    doc.add_paragraph(
        "Analiza jakości danych została wykonana za pomocą zapytań SQL w pliku sql/verification_km3.sql. "
        "Testy obejmowały liczbę rekordów, duplikaty na ziarnie faktów, zakresy wartości, spójność kluczy "
        "obcych oraz kompletność zakresu lat."
    )
    add_table(
        doc,
        ["Obiekt", "Liczba rekordów"],
        [
            ["staging.STG_COUNTRY", "40"],
            ["staging.STG_WEATHER", "84378"],
            ["staging.STG_ECONOMY", "429"],
            ["staging.STG_MENTAL_HEALTH", "2926"],
            ["dw.FACT_WEATHER", "84378"],
            ["dw.FACT_ECONOMY", "429"],
            ["dw.FACT_MENTAL_HEALTH", "2926"],
        ],
    )
    add_table(
        doc,
        ["Test", "Wynik"],
        [
            ["Duplikaty na ziarnie staging i faktów", "0 wykrytych problemów"],
            ["Wartości poza zakresem", "0 wykrytych problemów"],
            ["Brakujące powiązania faktów z wymiarami", "0 wykrytych problemów"],
            ["Zakres lat w faktach", "2000-2010"],
            ["Widoki agregujące pogodę", "AGG_WEATHER_MONTHLY: 2772; AGG_WEATHER_ANNUAL: 231; AGG_WEATHER_SEASONAL: 924"],
        ],
    )
    doc.add_paragraph(
        "Uwaga: podczas pobierania danych pogodowych Open-Meteo wystąpiły limity API. Zastosowano cache "
        "lokalny oraz możliwość kontynuowania procesu dla danych już pobranych. Finalny test potwierdził "
        "poprawność danych dostępnych w hurtowni i ich wykorzystanie w Power BI."
    )

    add_heading(doc, "9. Testy funkcjonalne")
    doc.add_paragraph(
        "Testy funkcjonalne potwierdziły działanie każdej warstwy rozwiązania: ETL, staging, hurtowni danych, "
        "warstwy BI oraz raportów. Szczegółowe scenariusze znajdują się w tests/functional_tests.md."
    )
    add_table(
        doc,
        ["ID", "Cel", "Oczekiwany wynik", "Potwierdzenie"],
        [
            ["TF-01", "Połączenie ETL z bazą", "Brak błędów logowania", "Log z konsoli"],
            ["TF-02", "Załadowanie staging", "Niezerowe liczby rekordów w STG_*", "01_ROW_COUNTS"],
            ["TF-03", "Załadowanie faktów", "Niezerowe liczby rekordów w FACT_*", "01_ROW_COUNTS"],
            ["TF-04", "Ponowne ładowanie", "Brak duplikatów", "02_DUPLICATE_GRAIN_CHECKS"],
            ["TF-05", "Transformacja pogody", "Poprawne zakresy i konwersje", "03_VALUE_RANGE_CHECKS"],
            ["TF-06", "Transformacja ekonomii", "Poprawne lata i zakresy wartości", "03_VALUE_RANGE_CHECKS, 06_DATE_RANGE_COVERAGE"],
            ["TF-07", "Transformacja zdrowia psychicznego", "Poprawne powiązania z DIM_DISORDER", "04_FOREIGN_KEY_COVERAGE"],
            ["TF-08", "Widoki agregujące", "Widoki zwracają rekordy", "08_AGGREGATE_VIEWS"],
            ["TF-09", "Power BI", "Raport odświeża się i pokazuje dane", "Power_BI_Raport.pdf"],
            ["TF-10", "End-to-end", "Dane przechodzą od źródeł do raportu", "09_END_TO_END_ANALYTIC_QUERY"],
        ],
    )
    doc.add_paragraph("[MIEJSCE NA ZRZUTY EKRANU TESTÓW SQL I POWER BI]")

    add_heading(doc, "10. Podsumowanie rezultatów")
    doc.add_paragraph(
        "Projekt zakończył się implementacją działającego systemu DW/BI. Dane z różnych źródeł zostały "
        "zintegrowane w SQL Serverze, poddane transformacjom i udostępnione w warstwie raportowej Power BI. "
        "Testy jakości danych nie wykazały duplikatów, wartości poza zakresem ani brakujących powiązań "
        "z wymiarami. System pozwala analizować wskaźniki zdrowia psychicznego w kontekście regionu, czasu, "
        "warunków klimatycznych i wybranych zmiennych ekonomicznych."
    )

    add_heading(doc, "11. Podział pracy w zespole")
    doc.add_paragraph("Sekcja do doprecyzowania przez zespół przed oddaniem. Proponowany układ:")
    add_table(
        doc,
        ["Osoba", "Zakres prac"],
        [
            ["Liwia Jankowska", "Źródła danych, przygotowanie danych CSV, część raportowa / dokumentacja - do doprecyzowania."],
            ["Wiktoria Kawa", "Model danych, STTM, walidacja danych / Power BI - do doprecyzowania."],
            ["Wojciech Łukaszyk", "SQL Server, ETL/ELT, testy funkcjonalne, integracja Power BI - do doprecyzowania."],
        ],
    )

    add_heading(doc, "12. Załączniki")
    add_bullets(
        doc,
        [
            "Power_BI_Raport.pdf - wyeksportowany raport Power BI.",
            "dw_climate_mental_health_DDL-3.sql - skrypt modelu hurtowni.",
            "sql/load_dw_from_staging.sql - transformacje ELT.",
            "sql/verification_km3.sql - zapytania testowe i weryfikacyjne.",
            "etl/ - kod źródłowy procesu ETL.",
            "STTM_KM2-1.xlsx - dokument STTM, do aktualizacji zgodnie z finalnym ETL.",
        ],
    )

    doc.save(OUT_DOCX)


def add_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    body.text = ""
    for idx, bullet in enumerate(bullets):
        paragraph = body.text_frame.paragraphs[0] if idx == 0 else body.text_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = PptPt(20)


def build_presentation() -> None:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Climate, Economy and Mental Health in Europe"
    slide.placeholders[1].text = "Raport finalny KM3 | DW/BI | 2000-2010"

    add_slide(
        prs,
        "Cel projektu",
        [
            "Budowa systemu DW/BI integrującego dane klimatyczne, ekonomiczne i zdrowia psychicznego.",
            "Zakres: kraje europejskie, lata 2000-2010.",
            "Odbiorca: analityk zdrowia publicznego lub decydent.",
        ],
    )
    add_slide(
        prs,
        "Architektura rozwiązania",
        [
            "Źródła: Open-Meteo API oraz lokalne CSV.",
            "ETL: Python, cache API, mapowania i zapis do staging.",
            "ELT/DW: SQL Server, schematy staging i dw.",
            "BI: Power BI Desktop, raporty interaktywne i eksport PDF.",
        ],
    )
    add_slide(
        prs,
        "Model hurtowni danych",
        [
            "Schemat konstelacji z trzema tabelami faktów.",
            "Fakty: FACT_WEATHER, FACT_ECONOMY, FACT_MENTAL_HEALTH.",
            "Wymiary: kraj, data, rok, miesiąc, zaburzenie.",
            "Widoki pogodowe: miesięczne, roczne i sezonowe.",
        ],
    )
    add_slide(
        prs,
        "Proces ETL/ELT",
        [
            "Python pobiera dane z API i lokalnych CSV.",
            "Dane trafiają do tabel staging.STG_*.",
            "SQL Server transformuje dane ze stagingu do modelu dw.",
            "Ponowne ładowanie jest idempotentne i nie tworzy duplikatów.",
        ],
    )
    add_slide(
        prs,
        "Jakość danych i testy",
        [
            "Brak duplikatów na ziarnie staging i faktów.",
            "Brak wartości poza zakresem w testach SQL.",
            "Brak brakujących powiązań faktów z wymiarami.",
            "Zakres faktów: lata 2000-2010.",
        ],
    )
    add_slide(
        prs,
        "Warstwa Power BI",
        [
            "Mental Health Map - przestrzenna analiza prevalencji.",
            "Mental Health & Economy - zależności z PKB i bezrobociem.",
            "Climate & Mental Health - nasłonecznienie, temperatura i zdrowie psychiczne.",
            "Alcohol & Unemployment - alkohol, bezrobocie i zaburzenia alkoholowe.",
            "What-if Simulation - scenariusz zmiany temperatury zimowej.",
        ],
    )
    add_slide(
        prs,
        "Najważniejsze wyniki",
        [
            "System umożliwia analizę danych według kraju, regionu, roku i zaburzenia.",
            "Agregaty pogodowe rozwiązują różnicę granulacji między pogodą dzienną a danymi rocznymi.",
            "Raporty Power BI potwierdzają działanie przepływu end-to-end.",
        ],
    )
    add_slide(
        prs,
        "Ograniczenia",
        [
            "Open-Meteo stosuje limity API, dlatego dodano lokalny cache i retry.",
            "Nie wszystkie pozycje słownika zaburzeń mają dane faktowe w źródłowym CSV.",
            "What-if jest symulacją scenariuszową, a nie modelem medycznym.",
        ],
    )
    add_slide(
        prs,
        "Podsumowanie",
        [
            "Zaimplementowano działający system DW/BI.",
            "Dane zostały zintegrowane, przetestowane i udostępnione w Power BI.",
            "Projekt spełnia wymagania KM3: ETL/ELT, hurtownia, BI, raporty i testy funkcjonalne.",
        ],
    )

    prs.save(OUT_PPTX)


if __name__ == "__main__":
    build_report()
    build_presentation()
    print(f"Created: {OUT_DOCX}")
    print(f"Created: {OUT_PPTX}")
